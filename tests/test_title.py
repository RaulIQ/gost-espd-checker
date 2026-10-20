from docx import Document
from hamcrest import assert_that, contains_string, empty, has_length

from gost_espd_checker.title import Title
from tests.rulebook import rulebook


def test_accepts_title_without_period() -> None:
    assert_that(
        Title(
            Document().add_paragraph("1.2. Функции программы"),
            "заголовок",
            rulebook(),
            hyphenation=False,
        ).violations(),
        empty(),
        "title without a period was not accepted",
    )


def test_accepts_two_sentences_without_final_period() -> None:
    assert_that(
        Title(
            Document().add_paragraph("3. НАСТРОЙКА. ПРОВЕРКА"),
            "заголовок",
            rulebook(),
            hyphenation=False,
        ).violations(),
        empty(),
        "two sentences without a final period were not accepted",
    )


def test_rejects_period_at_end() -> None:
    assert_that(
        str(
            Title(
                Document().add_paragraph("2.1. Состав программы."),
                "заголовок «2.1. Состав программы.»",
                rulebook(),
                hyphenation=False,
            ).violations()[0]
        ),
        contains_string(
            "заголовок «2.1. Состав программы.»: ошибка ESPD-106-12:"
            " в конце заголовка стоит точка"
        ),
        "period at the end of the title was not rejected",
    )


def test_rejects_soft_hyphen() -> None:
    assert_that(
        str(
            Title(
                Document().add_paragraph("4. НАСТРОЙКА ПОЛЬЗОВАТЕ\u00adЛЯ"),
                "заголовок",
                rulebook(),
                hyphenation=False,
            ).violations()[0]
        ),
        contains_string("ESPD-106-11: в заголовке стоит перенос"),
        "soft hyphen in the title was not rejected",
    )


def test_rejects_automatic_hyphenation_of_document() -> None:
    assert_that(
        str(
            Title(
                Document().add_paragraph("5. ДОПОЛНИТЕЛЬНЫЕ ВОЗМОЖНОСТИ"),
                "заголовок",
                rulebook(),
                hyphenation=True,
            ).violations()[0]
        ),
        contains_string(
            "ESPD-106-11: в заголовке разрешены автоматические переносы"
        ),
        "automatic hyphenation of the title was not rejected",
    )


def test_reports_period_and_hyphen_together() -> None:
    assert_that(
        Title(
            Document().add_paragraph("6.1. Сообще\u00adния."),
            "заголовок",
            rulebook(),
            hyphenation=False,
        ).violations(),
        has_length(2),
        "period and hyphen were not reported together",
    )
