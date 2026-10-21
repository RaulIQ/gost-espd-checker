from docx import Document
from docx.enum.text import WD_BREAK
from hamcrest import assert_that, equal_to

from gost_espd_checker.front import Front


def test_skips_approval_sheet() -> None:
    document = Document()
    document.add_paragraph("ЛИСТ УТВЕРЖДЕНИЯ").add_run().add_break(
        WD_BREAK.PAGE
    )
    document.add_paragraph("RU.17701729.00031-01 13 01")
    assert_that(
        int(Front(document)),
        equal_to(1),
        "title page after the approval sheet was not found",
    )


def test_points_past_pages_of_empty_document() -> None:
    assert_that(
        int(Front(Document())),
        equal_to(0),
        "empty document did not point past its pages",
    )
