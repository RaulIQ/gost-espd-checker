from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Mm
from hamcrest import assert_that, equal_to, is_

from gost_espd_checker.effective import Effective


def test_reads_alignment_of_paragraph() -> None:
    paragraph = Document().add_paragraph("Состав программы")
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    assert_that(
        Effective(paragraph).alignment(),
        equal_to(WD_PARAGRAPH_ALIGNMENT.RIGHT),
        "alignment of the paragraph was not read",
    )


def test_inherits_alignment_through_base_style() -> None:
    document = Document()
    document.styles[
        "Normal"
    ].paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    assert_that(
        Effective(
            document.add_paragraph("ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
        ).alignment(),
        equal_to(WD_PARAGRAPH_ALIGNMENT.CENTER),
        "alignment was not inherited through the base style",
    )


def test_aligns_left_without_any_alignment() -> None:
    assert_that(
        Effective(Document().add_paragraph("Связи программ")).alignment(),
        equal_to(WD_PARAGRAPH_ALIGNMENT.LEFT),
        "paragraph without alignment was not aligned left",
    )


def test_inherits_first_line_indent_from_style() -> None:
    document = Document()
    document.styles["Heading 2"].paragraph_format.first_line_indent = Mm(12.5)
    assert_that(
        Effective(
            document.add_paragraph("Установка программы", style="Heading 2")
        )
        .indent()
        .twips,
        equal_to(Mm(12.5).twips),
        "first line indent was not inherited from the style",
    )


def test_gives_zero_indent_without_any_indent() -> None:
    assert_that(
        Effective(Document().add_paragraph("Условия проверки")).indent(),
        equal_to(0),
        "paragraph without indent did not get a zero indent",
    )


def test_sees_capitals_set_by_style() -> None:
    document = Document()
    document.styles["Heading 1"].font.all_caps = True
    assert_that(
        Effective(
            document.add_paragraph("Настройка программы", style="Heading 1")
        ).capitals(),
        is_(True),
        "capitals set by the style were not seen",
    )


def test_denies_capitals_when_one_run_is_plain() -> None:
    paragraph = Document().add_paragraph()
    paragraph.add_run("ПРОВЕРКА ").font.all_caps = True
    paragraph.add_run("программы")
    assert_that(
        Effective(paragraph).capitals(),
        is_(False),
        "capitals were not denied for a paragraph with a plain run",
    )


def test_allows_hyphenation_by_default() -> None:
    assert_that(
        Effective(Document().add_paragraph("Контрольный пример")).hyphenation(),
        is_(True),
        "hyphenation was not allowed by default",
    )


def test_sees_hyphenation_suppressed_by_style() -> None:
    document = Document()
    document.styles["Heading 2"].element.get_or_add_pPr().append(
        parse_xml(f"<w:suppressAutoHyphens {nsdecls('w')}/>")
    )
    assert_that(
        Effective(
            document.add_paragraph("Сообщения журнала", style="Heading 2")
        ).hyphenation(),
        is_(False),
        "hyphenation suppressed by the style was not seen",
    )


def test_lets_paragraph_turn_hyphenation_back_on() -> None:
    document = Document()
    document.styles["Heading 2"].element.get_or_add_pPr().append(
        parse_xml(f"<w:suppressAutoHyphens {nsdecls('w')}/>")
    )
    paragraph = document.add_paragraph("Состав программы", style="Heading 2")
    paragraph.paragraph_format.element.get_or_add_pPr().append(
        parse_xml(f'<w:suppressAutoHyphens {nsdecls("w")} w:val="0"/>')
    )
    assert_that(
        Effective(paragraph).hyphenation(),
        is_(True),
        "paragraph did not turn hyphenation back on",
    )
