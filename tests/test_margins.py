import secrets

import pytest
from docx import Document
from docx.shared import Twips
from hamcrest import assert_that, contains_string, empty, has_length

from gost_espd_checker.margins import Margins
from tests.conforming import conforming


def test_accepts_margins_of_form_one() -> None:
    assert_that(
        Margins(
            conforming(Document().sections[0]), secrets.randbelow(9) + 1
        ).violations(),
        empty(),
        "margins of form 1 were not accepted",
    )


@pytest.mark.parametrize(
    "side", ["top_margin", "left_margin", "right_margin", "bottom_margin"]
)
def test_rejects_margin_off_by_one_twip(side: str) -> None:
    section = conforming(Document().sections[0])
    twips = {
        "top_margin": 1417,
        "left_margin": 1134,
        "right_margin": 567,
        "bottom_margin": 850,
    }[side]
    setattr(section, side, Twips(twips + secrets.choice((-1, 1))))
    assert_that(
        Margins(section, secrets.randbelow(9) + 1).violations(),
        has_length(1),
        f"{side} one twip off the norm was not rejected",
    )


def test_rejects_missing_margin() -> None:
    section = conforming(Document().sections[0])
    section.bottom_margin = None
    assert_that(
        str(Margins(section, secrets.randbelow(9) + 1).violations()[0]),
        contains_string("нижнее поле не задано вместо 15 мм"),
        "missing margin was not reported as unset",
    )


def test_names_section_side_and_width_of_wrong_margin() -> None:
    number = secrets.randbelow(90) + 2
    section = conforming(Document().sections[0])
    section.left_margin = Twips(1701)
    assert_that(
        str(Margins(section, number).violations()[0]),
        contains_string(f"раздел {number}: левое поле 30,0 мм вместо 20 мм"),
        "violation did not name the section, the side and the margin",
    )
