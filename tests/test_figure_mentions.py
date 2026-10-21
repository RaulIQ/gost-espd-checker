from docx import Document
from docx.document import Document as Word
from hamcrest import assert_that, contains_exactly, empty

from gost_espd_checker.figure_mentions import FigureMentions
from tests.conforming import endorsed
from tests.drawing import pictured
from tests.rulebook import rulebook


def mentioned(*texts: str) -> Word:
    document = endorsed(Document())
    for number in (1, 2):
        pictured(document)
        document.add_paragraph(f"Рис. {number}")
    for text in texts:
        document.add_paragraph(text)
    return document


def test_accepts_mentions_in_form_of_standard() -> None:
    assert_that(
        FigureMentions(
            mentioned(
                "Связи показаны на рис. 1.",
                "Нажать кнопку BOOT (рис. 2).",
                "Повторить запись (см. рис. 2).",
            ),
            rulebook(),
        ).violations(),
        empty(),
        "mentions in the form of the standard were not accepted",
    )


def test_rejects_mention_by_full_word() -> None:
    assert_that(
        [
            str(violation)
            for violation in FigureMentions(
                mentioned("Окно показано на рисунке 2."), rulebook()
            ).violations()
        ],
        contains_exactly(
            "абзац «Окно показано на рисунке 2.»: ошибка ESPD-106-19: ссылка"
            " «рисунке 2» вместо «рис. 2» (ГОСТ 19.106-78, п. 2.3.1)"
        ),
        "mention by the full word was not rejected",
    )


def test_rejects_mention_of_missing_illustration() -> None:
    assert_that(
        [
            str(violation)
            for violation in FigureMentions(
                mentioned("Журнал приведён на рис. 9."), rulebook()
            ).violations()
        ],
        contains_exactly(
            "абзац «Журнал приведён на рис. 9.»: ошибка ESPD-106-19: ссылка"
            " на рис. 9, которого нет в документе (ГОСТ 19.106-78, п. 2.3.1)"
        ),
        "mention of a missing illustration was not rejected",
    )


def test_rejects_repeated_mention_without_reminder() -> None:
    assert_that(
        [
            str(violation)
            for violation in FigureMentions(
                mentioned(
                    "Кнопки BOOT (рис. 1).", "Снова нажать BOOT (рис. 1)."
                ),
                rulebook(),
            ).violations()
        ],
        contains_exactly(
            "абзац «Снова нажать BOOT (рис. 1).»: ошибка ESPD-106-19:"
            " повторная ссылка на рис. 1 без «см.» (ГОСТ 19.106-78, п. 2.3.1)"
        ),
        "repeated mention without the reminder was not rejected",
    )
