from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Mm
from hamcrest import assert_that, contains_string, empty, has_length

from gost_espd_checker.indents import Indents
from tests.bold import bold
from tests.rulebook import rulebook


def test_accepts_indented_point() -> None:
    document = Document()
    document.add_paragraph(
        "2.1.1. Программа состоит из загрузчика."
    ).paragraph_format.first_line_indent = Mm(12.5)
    assert_that(
        Indents(document, rulebook()).violations(),
        empty(),
        "indented point was not accepted",
    )


def test_rejects_point_without_indent() -> None:
    document = Document()
    document.add_paragraph("2.1.2. Загрузчик проверяет образ.")
    assert_that(
        str(Indents(document, rulebook()).violations()[0]),
        contains_string(
            "пункт «2.1.2. Загрузчик проверяет образ.»: ошибка ESPD-106-07:"
            " пункт начинается без абзацного отступа"
        ),
        "point without an indent was not rejected",
    )


def test_rejects_enumeration_with_parenthesis_without_indent() -> None:
    document = Document()
    document.add_paragraph("1) настройка зарядного устройства;")
    assert_that(
        str(Indents(document, rulebook()).violations()[0]),
        contains_string(
            "ESPD-106-07: перечисление начинается без абзацного отступа"
        ),
        "enumeration with a parenthesis but no indent was not rejected",
    )


def test_rejects_enumeration_with_dash_without_indent() -> None:
    document = Document()
    document.add_paragraph("— передача телеметрии на пульт;")
    assert_that(
        Indents(document, rulebook()).violations(),
        has_length(1),
        "enumeration with a dash but no indent was not rejected",
    )


def test_rejects_word_list_item_without_indent() -> None:
    document = Document()
    paragraph = document.add_paragraph("управление двигателями;")
    paragraph.paragraph_format.element.get_or_add_pPr().append(
        parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/>'
            '<w:numId w:val="2"/></w:numPr>'
        )
    )
    assert_that(
        Indents(document, rulebook()).violations(),
        has_length(1),
        "Word list item without an indent was not rejected",
    )


def test_accepts_enumeration_indented_from_left() -> None:
    document = Document()
    document.add_paragraph(
        "2) прием команд управления;"
    ).paragraph_format.left_indent = Mm(12.5)
    assert_that(
        Indents(document, rulebook()).violations(),
        empty(),
        "enumeration indented from the left was not accepted",
    )


def test_ignores_plain_paragraph_without_indent() -> None:
    document = Document()
    document.add_paragraph("Программа прошивает контроллер.")
    assert_that(
        Indents(document, rulebook()).violations(),
        empty(),
        "plain paragraph without an indent was not ignored",
    )


def test_ignores_titles() -> None:
    document = Document()
    bold(document, "1.1. Назначение программы")
    assert_that(
        Indents(document, rulebook()).violations(),
        empty(),
        "title was not ignored",
    )


def test_ignores_list_item_without_letters_or_digits() -> None:
    document = Document()
    paragraph = document.add_paragraph(".")
    paragraph.paragraph_format.element.get_or_add_pPr().append(
        parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="4"/></w:numPr>')
    )
    assert_that(
        Indents(document, rulebook()).violations(),
        empty(),
        "list item without letters or digits was not ignored",
    )
