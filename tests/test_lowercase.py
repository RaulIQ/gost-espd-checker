from docx import Document
from docx.document import Document as Word
from hamcrest import assert_that, contains_exactly, empty

from gost_espd_checker.lowercase import Lowercase
from tests.rulebook import rulebook


def listed(*records: str) -> Word:
    document = Document()
    document.add_paragraph("СОДЕРЖАНИЕ")
    for record in records:
        document.add_paragraph(f"{record}\t5")
    return document


def test_accepts_names_in_lowercase_with_abbreviations() -> None:
    assert_that(
        Lowercase(
            listed("1. Общие сведения", "1.2. Обмен данными через ПК по USB"),
            rulebook(),
        ).violations(),
        empty(),
        "names in lowercase with abbreviations were not accepted",
    )


def test_rejects_name_in_capitals() -> None:
    assert_that(
        [
            str(violation)
            for violation in Lowercase(
                listed("2. ОПИСАНИЕ ПРОГРАММЫ"), rulebook()
            ).violations()
        ],
        contains_exactly(
            "содержание: ошибка ESPD-106-28: наименование «2. ОПИСАНИЕ"
            " ПРОГРАММЫ» записано прописными буквами (ГОСТ 19.106-78,"
            " п. 2.1.4)"
        ),
        "name in capitals was not rejected",
    )


def test_rejects_name_starting_with_small_letter() -> None:
    assert_that(
        [
            str(violation)
            for violation in Lowercase(
                listed("3.1. установка программы"), rulebook()
            ).violations()
        ],
        contains_exactly(
            "содержание: ошибка ESPD-106-28: первая буква наименования"
            " «3.1. установка программы» строчная (ГОСТ 19.106-78, п. 2.1.4)"
        ),
        "name starting with a small letter was not rejected",
    )


def test_skips_document_without_contents() -> None:
    document = Document()
    document.add_paragraph("4. ПРОВЕРКА ПРОГРАММЫ\t9")
    assert_that(
        Lowercase(document, rulebook()).violations(),
        empty(),
        "document without contents was not skipped",
    )
