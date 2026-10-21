from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from hamcrest import assert_that, equal_to

from gost_espd_checker.printout import Printout
from tests.boxed import boxed


def test_starts_new_page_after_page_break() -> None:
    document = Document()
    document.add_paragraph("Аннотация к программе").add_run().add_break(
        WD_BREAK.PAGE
    )
    document.add_paragraph("Перечень записей")
    assert_that(
        [[line.text for line in page] for page in Printout(document).pages()],
        equal_to([["Аннотация к программе"], ["Перечень записей"]]),
        "page break did not start a new page",
    )


def test_starts_new_page_before_paragraph_with_page_break_before() -> None:
    document = Document()
    document.add_paragraph("Листов 12")
    document.add_paragraph(
        "АННОТАЦИЯ"
    ).paragraph_format.page_break_before = True
    assert_that(
        len(Printout(document).pages()),
        equal_to(2),
        "page break before the paragraph did not start a new page",
    )


def test_drops_pages_without_text() -> None:
    document = Document()
    document.add_paragraph("Вводная часть").add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph(" ").add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("Заключение")
    assert_that(
        len(Printout(document).pages()),
        equal_to(2),
        "page without text was not dropped",
    )


def test_reads_paragraphs_inside_content_control() -> None:
    document = Document()
    document.element.body.xpath("./w:sectPr")[0].addprevious(
        parse_xml(
            f"<w:sdt {nsdecls('w')}><w:sdtContent><w:p><w:r>"
            "<w:t>1. Общие сведения\t4</w:t>"
            "</w:r></w:p></w:sdtContent></w:sdt>"
        )
    )
    assert_that(
        Printout(document).pages()[0][0].text,
        equal_to("1. Общие сведения\t4"),
        "paragraph inside the content control was not read",
    )


def test_skips_floating_table_in_margin() -> None:
    document = Document()
    document.element.body.xpath("./w:sectPr")[0].addprevious(
        parse_xml(
            f"<w:tbl {nsdecls('w')}><w:tblPr><w:tblpPr w:leftFromText="
            '"180"/></w:tblPr><w:tr><w:tc><w:p><w:r><w:t>Инв. № подл.'
            "</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
        )
    )
    document.add_paragraph("Программа")
    assert_that(
        [line.text for line in Printout(document).pages()[0]],
        equal_to(["Программа"]),
        "floating table in the margin was not skipped",
    )


def test_reads_text_box_once_after_its_paragraph() -> None:
    document = Document()
    boxed(document, "УТВЕРЖДЕНО", "ВЛТН.868600-01-ЛУ")
    document.add_paragraph("Программное обеспечение")
    assert_that(
        [line.text for line in Printout(document).pages()[0]],
        equal_to(
            ["", "УТВЕРЖДЕНО", "ВЛТН.868600-01-ЛУ", "Программное обеспечение"]
        ),
        "text box was not read once right after its paragraph",
    )


def test_skips_text_boxes_of_grouped_shapes() -> None:
    document = Document()
    document.add_paragraph().paragraph_format.element.append(
        parse_xml(
            "<w:r"
            ' xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' xmlns:wpg="http://schemas.microsoft.com/office/word/2010/'
            'wordprocessingGroup"'
            ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/'
            'wordprocessingShape">'
            "<w:drawing><wpg:wgp><wps:wsp><wps:txbx><w:txbxContent><w:p><w:r>"
            "<w:t>Инв. № подл.</w:t></w:r></w:p></w:txbxContent></wps:txbx>"
            "</wps:wsp></wpg:wgp></w:drawing></w:r>"
        )
    )
    document.add_paragraph("Листов 12")
    assert_that(
        [line.text for line in Printout(document).pages()[0]],
        equal_to(["", "Листов 12"]),
        "text box of grouped shapes was not skipped",
    )
