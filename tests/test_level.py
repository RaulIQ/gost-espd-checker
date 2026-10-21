from docx import Document
from hamcrest import assert_that, equal_to

from gost_espd_checker.level import Level
from tests.bold import bold


def test_takes_level_from_outline_of_styled_document() -> None:
    assert_that(
        int(
            Level(
                Document().add_paragraph("Связи программ", style="Heading 2"),
                _styled=True,
            )
        ),
        equal_to(2),
        "level was not taken from the outline of a styled document",
    )


def test_takes_level_from_number_of_bold_paragraph() -> None:
    assert_that(
        int(Level(bold(Document(), "3.4. Запись программы"), _styled=False)),
        equal_to(2),
        "level was not taken from the number of a bold paragraph",
    )


def test_gives_zero_to_plain_numbered_paragraph() -> None:
    assert_that(
        int(
            Level(
                Document().add_paragraph("1. Подключить плату"), _styled=False
            )
        ),
        equal_to(0),
        "plain numbered paragraph was not left without a level",
    )
