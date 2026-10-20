import io
from pathlib import Path

from docx import Document
from docx.shared import Mm
from hamcrest import assert_that, equal_to

from gost_espd_checker.cli import Cli
from tests.conforming import conforming
from tests.rulebook import rulebook


def test_exits_with_zero_for_conforming_document(tmp_path: Path) -> None:
    document = Document()
    conforming(document.sections[0])
    path = tmp_path / "доклад.docx"
    document.save(str(path))
    assert_that(
        Cli(path, rulebook(), io.StringIO()).run(),
        equal_to(0),
        "conforming document did not exit with zero",
    )


def test_exits_with_one_for_wrong_document(tmp_path: Path) -> None:
    document = Document()
    conforming(document.sections[0]).right_margin = Mm(15)
    path = tmp_path / "отчёт.docx"
    document.save(str(path))
    assert_that(
        Cli(path, rulebook(), io.StringIO()).run(),
        equal_to(1),
        "wrong document did not exit with one",
    )


def test_prints_violation_of_wrong_document(tmp_path: Path) -> None:
    document = Document()
    conforming(document.sections[0]).right_margin = Mm(15)
    path = tmp_path / "отчёт.docx"
    document.save(str(path))
    out = io.StringIO()
    Cli(path, rulebook(), out).run()
    assert_that(
        out.getvalue(),
        equal_to(
            "раздел 1: ошибка ESPD-106-02: правое поле 15,0 мм вместо 10 мм"
            " (ГОСТ 19.106-78, п. 2.1.2)\n"
        ),
        "violation of the wrong document was not printed",
    )
