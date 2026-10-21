from docx import Document
from hamcrest import assert_that, contains_exactly

from gost_espd_checker.entries import Entries
from tests.bold import bold


def test_lists_titles_and_points_with_numbers() -> None:
    document = Document()
    bold(document, "2. НАЗНАЧЕНИЕ ПРОГРАММЫ")
    document.add_paragraph("2.1.4. Модуль связи передаёт данные.")
    document.add_paragraph("Программа состоит из модулей.")
    assert_that(
        [
            (level, number, dotted)
            for _, level, number, dotted in Entries(document)
        ],
        contains_exactly((1, (2,), True), (0, (2, 1, 4), True)),
        "titles and points did not come with their numbers",
    )
