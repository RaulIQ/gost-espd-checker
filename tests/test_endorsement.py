from docx import Document
from docx.document import Document as Word
from docx.enum.text import WD_BREAK, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from hamcrest import assert_that, contains_string, empty, has_length

from gost_espd_checker.endorsement import Endorsement
from tests.rulebook import rulebook


def titled(document: Word, *lines: str) -> Word:
    for line in lines:
        document.add_paragraph(line)
    return document


def test_accepts_inscription_in_top_left_corner() -> None:
    document = titled(
        Document(),
        "Утвержден",
        "RU.17701729.00001-01 32 01-ЛУ",
        "РУКОВОДСТВО СИСТЕМНОГО ПРОГРАММИСТА",
    )
    assert_that(
        Endorsement(document, rulebook()).violations(),
        empty(),
        "inscription in the top left corner was not accepted",
    )


def test_accepts_inscription_in_capitals() -> None:
    document = titled(Document(), "УТВЕРЖДЕН", "RU.04838275.00317-02 34 01-ЛУ")
    assert_that(
        Endorsement(document, rulebook()).violations(),
        empty(),
        "inscription in capitals was not accepted",
    )


def test_rejects_title_page_without_inscription() -> None:
    document = titled(
        Document(), "Программное обеспечение", "ВЛТН.868600-01 32 01"
    )
    assert_that(
        str(Endorsement(document, rulebook()).violations()[0]),
        contains_string(
            "титульный лист: ошибка ESPD-104-01:"
            " на титульном листе нет надписи «Утвержден»"
        ),
        "title page without the inscription was not rejected",
    )


def test_rejects_other_form_of_inscription() -> None:
    document = titled(Document(), "УТВЕРЖДЕНО", "RU.BMS.00001-01 32 01-ЛУ")
    assert_that(
        str(Endorsement(document, rulebook()).violations()[0]),
        contains_string("надпись «УТВЕРЖДЕНО» вместо «Утвержден»"),
        "other form of the inscription was not rejected",
    )


def test_rejects_centered_inscription() -> None:
    document = titled(Document(), "Утвержден", "RU.BMS.00001-01 32 01-ЛУ")
    document.paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    assert_that(
        str(Endorsement(document, rulebook()).violations()[0]),
        contains_string("надпись «Утвержден» стоит не в левом верхнем углу"),
        "centered inscription was not rejected",
    )


def test_rejects_inscription_below_other_text() -> None:
    document = titled(
        Document(),
        "Программное обеспечение",
        "Утвержден",
        "RU.BMS.00001-01 32 01-ЛУ",
    )
    assert_that(
        Endorsement(document, rulebook()).violations(),
        has_length(1),
        "inscription below other text was not rejected",
    )


def test_rejects_inscription_without_designation_of_approval_sheet() -> None:
    document = titled(
        Document(), "Утвержден", "РУКОВОДСТВО СИСТЕМНОГО ПРОГРАММИСТА"
    )
    assert_that(
        str(Endorsement(document, rulebook()).violations()[0]),
        contains_string(
            "под надписью «Утвержден» нет обозначения листа утверждения"
        ),
        "inscription without the designation was not rejected",
    )


def test_skips_approval_sheet_before_title_page() -> None:
    document = titled(Document(), "СОГЛАСОВАНО", "ЛИСТ УТВЕРЖДЕНИЯ")
    document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    titled(document, "Утвержден", "RU.BMS.00001-01 32 01-ЛУ")
    assert_that(
        Endorsement(document, rulebook()).violations(),
        empty(),
        "approval sheet before the title page was not skipped",
    )


def test_looks_only_at_first_page() -> None:
    document = titled(Document(), "Программное обеспечение")
    document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    titled(document, "Утвержден", "RU.BMS.00001-01 32 01-ЛУ")
    assert_that(
        Endorsement(document, rulebook()).violations(),
        has_length(1),
        "inscription on the second page was not ignored",
    )


def test_ignores_floating_table_in_margin() -> None:
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Инв. № подл."
    document.element.body.xpath("./w:tbl/w:tblPr")[0].append(
        parse_xml(
            f'<w:tblpPr {nsdecls("w")} w:horzAnchor="page"'
            ' w:vertAnchor="page" w:tblpX="397" w:tblpY="8334"/>'
        )
    )
    titled(document, "Утвержден", "RU.BMS.00001-01 32 01-ЛУ")
    assert_that(
        Endorsement(document, rulebook()).violations(),
        empty(),
        "floating table in the margin was not ignored",
    )


def test_reads_inscription_inside_table() -> None:
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Утвержден"
    table.cell(1, 0).text = "RU.BMS.00001-01 32 01-ЛУ"
    assert_that(
        Endorsement(document, rulebook()).violations(),
        empty(),
        "inscription inside a table was not read",
    )


def test_rejects_empty_document() -> None:
    assert_that(
        Endorsement(Document(), rulebook()).violations(),
        has_length(1),
        "empty document was not rejected for the missing inscription",
    )
