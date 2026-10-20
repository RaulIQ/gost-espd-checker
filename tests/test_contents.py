from docx import Document
from docx.document import Document as Word
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from hamcrest import assert_that, contains_exactly, empty, has_item

from gost_espd_checker.contents import Contents
from tests.rulebook import rulebook


def paged(document: Word, *pages: tuple[str, ...]) -> Word:
    for page in pages:
        for line in page:
            document.add_paragraph(line)
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    return document


def test_accepts_contents_on_own_page_after_annotation() -> None:
    document = paged(
        Document(),
        ("АННОТАЦИЯ", "Документ описывает настройку программы."),
        ("СОДЕРЖАНИЕ", "1. Назначение программы\t4", "1.1. Функции\t4"),
        ("1. НАЗНАЧЕНИЕ ПРОГРАММЫ",),
    )
    assert_that(
        Contents(document, rulebook()).violations(),
        empty(),
        "contents on its own page after the annotation was not accepted",
    )


def test_skips_document_without_contents() -> None:
    document = paged(Document(), ("ВВЕДЕНИЕ", "Основания для разработки."))
    assert_that(
        Contents(document, rulebook()).violations(),
        empty(),
        "document without contents was not skipped",
    )


def test_rejects_other_title_of_contents() -> None:
    document = paged(Document(), ("Оглавление", "2. Структура\t7"))
    assert_that(
        [
            str(violation)
            for violation in Contents(document, rulebook()).violations()
        ],
        contains_exactly(
            "содержание: ошибка ESPD-106-16: заголовок «Оглавление» вместо"
            " «СОДЕРЖАНИЕ» (ГОСТ 19.106-78, п. 2.1.4)"
        ),
        "other title of the contents was not rejected",
    )


def test_rejects_contents_numbered_as_section() -> None:
    document = paged(Document(), ("1. СОДЕРЖАНИЕ", "2. Условия выполнения\t6"))
    assert_that(
        [
            str(violation)
            for violation in Contents(document, rulebook()).violations()
        ],
        contains_exactly(
            "содержание: ошибка ESPD-106-16: содержание пронумеровано как"
            " раздел (ГОСТ 19.106-78, п. 2.1.4)"
        ),
        "contents numbered as a section was not rejected",
    )


def test_rejects_text_before_contents_on_same_page() -> None:
    document = paged(
        Document(),
        ("АННОТАЦИЯ",),
        ("Оформление выполнено по ЕСПД.", "СОДЕРЖАНИЕ", "3. Настройка\t8"),
    )
    assert_that(
        [
            str(violation)
            for violation in Contents(document, rulebook()).violations()
        ],
        contains_exactly(
            "содержание: ошибка ESPD-106-16: перед заголовком «СОДЕРЖАНИЕ»"
            " на той же странице есть текст (ГОСТ 19.106-78, п. 2.1.4)"
        ),
        "text before the contents on the same page was not rejected",
    )


def test_rejects_contents_before_annotation() -> None:
    document = paged(
        Document(),
        ("СОДЕРЖАНИЕ", "4. Проверка программы\t9"),
        ("АННОТАЦИЯ", "Документ предназначен для персонала."),
    )
    assert_that(
        [
            str(violation)
            for violation in Contents(document, rulebook()).violations()
        ],
        contains_exactly(
            "содержание: ошибка ESPD-106-16: содержание стоит перед"
            " аннотацией (ГОСТ 19.106-78, п. 2.1.4)"
        ),
        "contents before the annotation was not rejected",
    )


def test_rejects_annotation_listed_in_contents() -> None:
    document = paged(
        Document(),
        ("АННОТАЦИЯ",),
        ("СОДЕРЖАНИЕ", "Аннотация\t2", "1. Общие сведения\t4"),
    )
    assert_that(
        [
            str(violation)
            for violation in Contents(document, rulebook()).violations()
        ],
        contains_exactly(
            "содержание: ошибка ESPD-106-16: в содержание включена строка"
            " «Аннотация» (ГОСТ 19.106-78, п. 2.1.4)"
        ),
        "annotation listed in the contents was not rejected",
    )


def test_rejects_contents_listed_in_itself() -> None:
    document = paged(
        Document(),
        ("АННОТАЦИЯ",),
        ("СОДЕРЖАНИЕ", "СОДЕРЖАНИЕ\t3", "5. Сообщения оператору\t11"),
    )
    assert_that(
        [
            str(violation)
            for violation in Contents(document, rulebook()).violations()
        ],
        has_item(
            "содержание: ошибка ESPD-106-16: в содержание включена строка"
            " «СОДЕРЖАНИЕ» (ГОСТ 19.106-78, п. 2.1.4)"
        ),
        "contents listed in itself was not rejected",
    )


def test_reads_lines_inside_table_of_contents_field() -> None:
    document = paged(Document(), ("АННОТАЦИЯ",), ("СОДЕРЖАНИЕ",))
    document.element.body.xpath("./w:sectPr")[0].addprevious(
        parse_xml(
            f"<w:sdt {nsdecls('w')}><w:sdtContent><w:p><w:r>"
            "<w:t>АННОТАЦИЯ\t2</w:t>"
            "</w:r></w:p></w:sdtContent></w:sdt>"
        )
    )
    assert_that(
        [
            str(violation)
            for violation in Contents(document, rulebook()).violations()
        ],
        contains_exactly(
            "содержание: ошибка ESPD-106-16: в содержание включена строка"
            " «АННОТАЦИЯ» (ГОСТ 19.106-78, п. 2.1.4)"
        ),
        "line inside the table of contents field was not read",
    )
