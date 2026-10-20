import secrets

from docx import Document
from docx.shared import Mm, Twips
from hamcrest import assert_that, contains_string, empty, has_length

from gost_espd_checker.page_size import PageSize
from tests.conforming import conforming
from tests.rulebook import rulebook


def test_accepts_a4_portrait_sheet() -> None:
    section = conforming(Document().sections[0])
    assert_that(
        PageSize(section, secrets.randbelow(9) + 1, rulebook()).violations(),
        empty(),
        "A4 portrait sheet was not accepted",
    )


def test_accepts_a3_landscape_sheet() -> None:
    section = conforming(Document().sections[0])
    section.page_width = Mm(420)
    assert_that(
        PageSize(section, secrets.randbelow(9) + 1, rulebook()).violations(),
        empty(),
        "A3 landscape sheet was not accepted",
    )


def test_rejects_a4_landscape_sheet() -> None:
    section = conforming(Document().sections[0])
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    assert_that(
        PageSize(section, secrets.randbelow(9) + 1, rulebook()).violations(),
        has_length(1),
        "A4 landscape sheet was not rejected",
    )


def test_rejects_sheet_one_twip_wider_than_a4() -> None:
    section = conforming(Document().sections[0])
    section.page_width = Twips(Mm(210).twips + 1)
    assert_that(
        PageSize(section, secrets.randbelow(9) + 1, rulebook()).violations(),
        has_length(1),
        "sheet one twip wider than A4 was not rejected",
    )


def test_rejects_sheet_without_width() -> None:
    section = conforming(Document().sections[0])
    section.page_width = None
    assert_that(
        PageSize(section, secrets.randbelow(9) + 1, rulebook()).violations(),
        has_length(1),
        "sheet without width was not rejected",
    )


def test_names_section_and_size_of_wrong_sheet() -> None:
    number = secrets.randbelow(90) + 2
    section = conforming(Document().sections[0])
    section.page_width = Mm(215.9)
    section.page_height = Mm(279.4)
    assert_that(
        str(PageSize(section, number, rulebook()).violations()[0]),
        contains_string(
            f"раздел {number}: ошибка ESPD-106-01: лист 215,9 мм на 279,4 мм"
        ),
        "violation did not name the section and the size of the sheet",
    )
