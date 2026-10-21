from docx import Document
from hamcrest import assert_that, contains_exactly, equal_to

from gost_espd_checker.lineage import Lineage


def test_puts_paragraph_before_its_styles() -> None:
    document = Document()
    paragraph = document.add_paragraph("Состав программы", style="Heading 3")
    assert_that(
        Lineage(paragraph).owners(),
        contains_exactly(
            paragraph.paragraph_format.element,
            document.styles["Heading 3"].element,
            document.styles["Normal"].element,
        ),
        "paragraph and its styles did not come in order of inheritance",
    )


def test_leaves_out_runs_without_text() -> None:
    paragraph = Document().add_paragraph("Связи программ")
    paragraph.add_run("   ")
    assert_that(
        len(Lineage(paragraph).flags(lambda font: font.italic)),
        equal_to(1),
        "run without text was not left out",
    )


def test_lets_run_override_its_style() -> None:
    document = Document()
    document.styles["Heading 2"].font.bold = True
    document.add_paragraph(style="Heading 2").add_run(
        "Установка программы"
    ).bold = False
    assert_that(
        [
            "жирный" if flag else "обычный"
            for flag in Lineage(document.paragraphs[0]).flags(
                lambda font: font.bold
            )
        ],
        contains_exactly("обычный"),
        "run did not override the bold of its style",
    )


def test_takes_value_from_style_when_run_sets_none() -> None:
    document = Document()
    document.styles["Heading 1"].font.all_caps = True
    assert_that(
        [
            "прописные" if flag else "строчные"
            for flag in Lineage(
                document.add_paragraph("Общие сведения", style="Heading 1")
            ).flags(lambda font: font.all_caps)
        ],
        contains_exactly("прописные"),
        "value of the style was not taken for the run",
    )
