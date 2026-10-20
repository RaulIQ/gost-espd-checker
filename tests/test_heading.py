import secrets

from docx import Document
from hamcrest import assert_that, contains_string, empty, has_length

from gost_espd_checker.heading import Heading
from tests.conforming import conforming, headed


def test_accepts_page_field_above_designation() -> None:
    assert_that(
        Heading(
            conforming(Document().sections[0]), secrets.randbelow(9) + 1
        ).violations(),
        empty(),
        "page field above the designation was not accepted",
    )


def test_accepts_complex_page_field() -> None:
    section = headed(
        Document().sections[0],
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> PAGE  \\* MERGEFORMAT'
        "</w:instrText></w:r>"
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        "<w:r><w:t>4</w:t></w:r>"
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
    )
    headed(section, "<w:r><w:t>RU.04838275.00317-02 51 01-1</w:t></w:r>")
    assert_that(
        Heading(section, secrets.randbelow(9) + 1).violations(),
        empty(),
        "complex page field was not accepted",
    )


def test_rejects_header_without_page_number() -> None:
    number = secrets.randbelow(90) + 2
    section = headed(
        Document().sections[0],
        "<w:r><w:t>RU.17701729.00001-01 34 01</w:t></w:r>",
    )
    assert_that(
        str(Heading(section, number).violations()[0]),
        contains_string(
            f"раздел {number}: в верхнем колонтитуле нет номера страницы"
        ),
        "header without a page number was not rejected",
    )


def test_rejects_page_count_as_page_number() -> None:
    section = headed(
        Document().sections[0],
        '<w:fldSimple w:instr=" NUMPAGES ">'
        "<w:r><w:t>9</w:t></w:r></w:fldSimple>",
    )
    headed(section, "<w:r><w:t>RU.17701729.00001-01 34 01</w:t></w:r>")
    assert_that(
        str(Heading(section, secrets.randbelow(9) + 1).violations()[0]),
        contains_string("нет номера страницы"),
        "page count was not rejected as a page number",
    )


def test_rejects_header_without_designation() -> None:
    section = headed(
        Document().sections[0],
        '<w:fldSimple w:instr=" PAGE "><w:r><w:t>3</w:t></w:r></w:fldSimple>',
    )
    assert_that(
        str(Heading(section, secrets.randbelow(9) + 1).violations()[0]),
        contains_string("нет обозначения документа под номером страницы"),
        "header without a designation was not rejected",
    )


def test_rejects_designation_above_page_number() -> None:
    section = headed(
        Document().sections[0],
        "<w:r><w:t>RU.17701729.00001-01 34 01</w:t></w:r>",
    )
    headed(
        section,
        '<w:fldSimple w:instr=" PAGE "><w:r><w:t>5</w:t></w:r></w:fldSimple>',
    )
    assert_that(
        Heading(section, secrets.randbelow(9) + 1).violations(),
        has_length(1),
        "designation above the page number was not rejected",
    )


def test_ignores_header_of_first_page() -> None:
    section = conforming(Document().sections[0])
    section.different_first_page_header_footer = True
    assert_that(
        Heading(section, secrets.randbelow(9) + 1).violations(),
        empty(),
        "empty header of the title page was not ignored",
    )


def test_reads_header_linked_to_previous_section() -> None:
    document = Document()
    conforming(document.sections[0])
    assert_that(
        Heading(document.add_section(), 2).violations(),
        empty(),
        "header linked to the previous section was not read",
    )
