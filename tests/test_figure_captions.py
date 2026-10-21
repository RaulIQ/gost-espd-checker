from docx import Document
from hamcrest import assert_that, contains_exactly, empty

from gost_espd_checker.figure_captions import FigureCaptions
from tests.conforming import endorsed
from tests.drawing import pictured
from tests.rulebook import rulebook


def test_accepts_title_over_and_number_under() -> None:
    document = endorsed(Document())
    document.add_paragraph("Окно PuTTY Configuration")
    pictured(document)
    document.add_paragraph("Тип соединения — Serial.")
    document.add_paragraph("Рис. 6")
    assert_that(
        FigureCaptions(document, rulebook()).violations(),
        empty(),
        "title over and number under the illustration were not accepted",
    )


def test_rejects_number_over_illustration() -> None:
    document = endorsed(Document())
    document.add_paragraph("Рис. 5")
    pictured(document)
    assert_that(
        [
            str(violation)
            for violation in FigureCaptions(document, rulebook()).violations()
        ],
        contains_exactly(
            "иллюстрация «Рис. 5»: ошибка ESPD-106-20: номер стоит над"
            " иллюстрацией, хотя должен быть под ней (ГОСТ 19.106-78, п. 2.3.1)"
        ),
        "number over the illustration was not rejected",
    )


def test_rejects_title_next_to_number() -> None:
    document = endorsed(Document())
    pictured(document)
    document.add_paragraph("Рис. 2 — Расположение кнопок")
    assert_that(
        [
            str(violation)
            for violation in FigureCaptions(document, rulebook()).violations()
        ],
        contains_exactly(
            "иллюстрация «Рис. 2 — Расположение кнопок»: ошибка"
            " ESPD-106-20: наименование записано в строке номера, хотя"
            " должно быть над иллюстрацией (ГОСТ 19.106-78, п. 2.3.1)"
        ),
        "title next to the number was not rejected",
    )
