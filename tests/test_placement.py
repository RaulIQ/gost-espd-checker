from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from hamcrest import assert_that, equal_to, is_

from gost_espd_checker.lineage import Lineage
from gost_espd_checker.placement import Placement


def test_takes_outline_level_from_style() -> None:
    assert_that(
        Placement(
            Lineage(
                Document().add_paragraph("Связи программ", style="Heading 2")
            )
        ).outline(),
        equal_to(1),
        "outline level was not taken from the style",
    )


def test_takes_outline_level_from_paragraph() -> None:
    paragraph = Document().add_paragraph("ОБЩИЕ СВЕДЕНИЯ")
    paragraph.paragraph_format.element.get_or_add_pPr().append(
        parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>')
    )
    assert_that(
        Placement(Lineage(paragraph)).outline(),
        equal_to(0),
        "outline level was not taken from the paragraph",
    )


def test_gives_body_text_level_without_outline() -> None:
    assert_that(
        Placement(
            Lineage(Document().add_paragraph("Текст программы"))
        ).outline(),
        equal_to(9),
        "paragraph without an outline level was not treated as body text",
    )


def test_sees_paragraph_in_word_list() -> None:
    paragraph = Document().add_paragraph("управление двигателями")
    paragraph.paragraph_format.element.get_or_add_pPr().append(
        parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/>'
            '<w:numId w:val="3"/></w:numPr>'
        )
    )
    assert_that(
        Placement(Lineage(paragraph)).listed(),
        is_(True),
        "paragraph in a Word list was not seen",
    )


def test_sees_list_set_by_style() -> None:
    document = Document()
    document.styles["Normal"].element.get_or_add_pPr().append(
        parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="5"/></w:numPr>')
    )
    assert_that(
        Placement(
            Lineage(document.add_paragraph("чтение данных датчиков"))
        ).listed(),
        is_(True),
        "list set by the style was not seen",
    )


def test_denies_list_removed_by_zero_number() -> None:
    document = Document()
    document.styles["Normal"].element.get_or_add_pPr().append(
        parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="5"/></w:numPr>')
    )
    paragraph = document.add_paragraph("Программа прошивает контроллер.")
    paragraph.paragraph_format.element.get_or_add_pPr().append(
        parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="0"/></w:numPr>')
    )
    assert_that(
        Placement(Lineage(paragraph)).listed(),
        is_(False),
        "list removed by the zero number was not denied",
    )


def test_denies_list_to_plain_paragraph() -> None:
    assert_that(
        Placement(
            Lineage(Document().add_paragraph("Текст программы"))
        ).listed(),
        is_(False),
        "list was not denied to a plain paragraph",
    )
