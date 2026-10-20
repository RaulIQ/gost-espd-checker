from docx import Document
from hamcrest import assert_that, is_

from gost_espd_checker.lettering import Lettering
from gost_espd_checker.lineage import Lineage


def test_sees_capitals_set_by_style() -> None:
    document = Document()
    document.styles["Heading 1"].font.all_caps = True
    assert_that(
        Lettering(
            Lineage(
                document.add_paragraph("Настройка программы", style="Heading 1")
            )
        ).capitals(),
        is_(True),
        "capitals set by the style were not seen",
    )


def test_denies_capitals_when_one_run_is_plain() -> None:
    paragraph = Document().add_paragraph()
    paragraph.add_run("ПРОВЕРКА ").font.all_caps = True
    paragraph.add_run("программы")
    assert_that(
        Lettering(Lineage(paragraph)).capitals(),
        is_(False),
        "capitals were not denied for a paragraph with a plain run",
    )


def test_sees_bold_set_by_style() -> None:
    document = Document()
    document.styles["Heading 2"].font.bold = True
    assert_that(
        Lettering(
            Lineage(
                document.add_paragraph("Состав программы", style="Heading 2")
            )
        ).bold(),
        is_(True),
        "bold set by the style was not seen",
    )


def test_denies_bold_when_only_some_words_are_bold() -> None:
    paragraph = Document().add_paragraph()
    paragraph.add_run("7. Перейти на вкладку ")
    paragraph.add_run("«Firmware»").bold = True
    assert_that(
        Lettering(Lineage(paragraph)).bold(),
        is_(False),
        "bold was not denied for a paragraph with some plain words",
    )


def test_denies_bold_to_paragraph_without_text() -> None:
    assert_that(
        Lettering(Lineage(Document().add_paragraph("  "))).bold(),
        is_(False),
        "bold was not denied to a paragraph without text",
    )
