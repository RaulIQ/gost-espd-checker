from docx import Document
from hamcrest import assert_that, equal_to

from gost_espd_checker.kind import Kind


def test_names_point_by_its_number() -> None:
    assert_that(
        str(Kind(Document().add_paragraph("2.3.1. Программа читает файл"))),
        equal_to("пункт"),
        "paragraph with a point number was not named a point",
    )


def test_names_enumeration_by_its_dash() -> None:
    assert_that(
        str(Kind(Document().add_paragraph("— кабель USB Type-C"))),
        equal_to("перечисление"),
        "paragraph with a dash was not named an enumeration",
    )


def test_leaves_plain_paragraph_unnamed() -> None:
    assert_that(
        str(Kind(Document().add_paragraph("Программа хранит журнал."))),
        equal_to(""),
        "plain paragraph did not stay unnamed",
    )
