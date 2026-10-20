from docx import Document
from hamcrest import assert_that, equal_to

from gost_espd_checker.headings import Headings
from tests.bold import bold


def test_takes_levels_from_outline() -> None:
    document = Document()
    document.add_paragraph("2. ОПИСАНИЕ ПРОГРАММЫ", style="Heading 1")
    document.add_paragraph("2.1. Состав программы", style="Heading 2")
    document.add_paragraph("Программа состоит из двух модулей.")
    assert_that(
        [level for _, level in Headings(document).levels()],
        equal_to([1, 2, 0]),
        "levels were not taken from the outline",
    )


def test_takes_levels_from_bold_numbers_without_outline() -> None:
    document = Document()
    bold(document, "4. ПРОВЕРКА ПРОГРАММЫ")
    bold(document, "4.2. Контрольный пример")
    document.add_paragraph("1. Подключить контроллер;")
    assert_that(
        [level for _, level in Headings(document).levels()],
        equal_to([1, 2, 0]),
        "levels were not taken from bold numbers",
    )


def test_leaves_out_lines_of_contents() -> None:
    document = Document()
    document.add_paragraph("1.1. Назначение программы\t4")
    document.add_paragraph("Программа прошивает контроллер.")
    assert_that(
        [paragraph.text for paragraph, _ in Headings(document).levels()],
        equal_to(["Программа прошивает контроллер."]),
        "lines of the contents were not left out",
    )
