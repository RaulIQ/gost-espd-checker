from docx import Document
from hamcrest import assert_that, contains_exactly, empty

from gost_espd_checker.records import Records


def test_ends_records_at_first_line_of_other_kind() -> None:
    document = Document()
    document.add_paragraph("Перечень сокращений\t23")
    document.add_paragraph("СОДЕРЖАНИЕ")
    document.add_paragraph("1. Назначение программы\t4")
    document.add_paragraph("Программа работает автономно.")
    document.add_paragraph("Приложение 2\t19")
    assert_that(
        list(Records(document)),
        contains_exactly("1. Назначение программы"),
        "records did not end at the first line of other kind",
    )


def test_lists_records_after_contents_title() -> None:
    document = Document()
    document.add_paragraph("Перечень сокращений\t23")
    document.add_paragraph("СОДЕРЖАНИЕ")
    document.add_paragraph("1. Назначение программы\t4")
    document.add_paragraph("Приложение 2\t19")
    assert_that(
        list(Records(document)),
        contains_exactly("1. Назначение программы", "Приложение 2"),
        "records after the contents title were not listed",
    )


def test_lists_nothing_without_contents() -> None:
    document = Document()
    document.add_paragraph("3.2. Запуск программы\t8")
    assert_that(
        list(Records(document)),
        empty(),
        "document without contents did not give an empty list",
    )
