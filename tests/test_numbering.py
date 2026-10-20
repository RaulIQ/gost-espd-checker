from docx import Document
from hamcrest import assert_that, contains_string, empty, has_length

from gost_espd_checker.numbering import Numbering
from tests.bold import bold
from tests.rulebook import rulebook


def test_accepts_sections_and_subsections_in_order() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    document.add_paragraph("1.1. Назначение", style="Heading 2")
    document.add_paragraph("1.2. Функции", style="Heading 2")
    document.add_paragraph("2. ОПИСАНИЕ", style="Heading 1")
    document.add_paragraph("2.1. Состав", style="Heading 2")
    assert_that(
        Numbering(document, rulebook()).violations(),
        empty(),
        "sections and subsections in order were not accepted",
    )


def test_rejects_skipped_section_number() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    document.add_paragraph("3. НАСТРОЙКА", style="Heading 1")
    assert_that(
        str(Numbering(document, rulebook()).violations()[0]),
        contains_string(
            "заголовок «3. НАСТРОЙКА»: ошибка ESPD-106-17:"
            " номер 3 не продолжает нумерацию после 1"
        ),
        "skipped section number was not rejected",
    )


def test_rejects_first_section_other_than_one() -> None:
    document = Document()
    document.add_paragraph("2. ОПИСАНИЕ", style="Heading 1")
    assert_that(
        str(Numbering(document, rulebook()).violations()[0]),
        contains_string(
            "заголовок «2. ОПИСАНИЕ»: ошибка ESPD-106-17: номер 2 вместо 1"
        ),
        "first section other than one was not rejected",
    )


def test_rejects_subsection_of_another_section() -> None:
    document = Document()
    bold(document, "1. ОБЩИЕ СВЕДЕНИЯ")
    bold(document, "2.1. Состав программы")
    assert_that(
        str(Numbering(document, rulebook()).violations()[0]),
        contains_string("номер 2.1 не продолжает нумерацию после 1"),
        "subsection of another section was not rejected",
    )


def test_rejects_skipped_subsection_number() -> None:
    document = Document()
    document.add_paragraph("4. ПРОВЕРКА", style="Heading 1")
    document.add_paragraph("4.1. Условия", style="Heading 2")
    document.add_paragraph("4.3. Заключение", style="Heading 2")
    assert_that(
        Numbering(document, rulebook()).violations(),
        has_length(2),
        "skipped subsection number was not rejected",
    )


def test_accepts_points_inside_subsection() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    document.add_paragraph("1.1. Назначение", style="Heading 2")
    document.add_paragraph("1.1.1. Программа прошивает контроллер.")
    document.add_paragraph("1.1.2. Программа читает журнал.")
    document.add_paragraph("1.2. Функции", style="Heading 2")
    assert_that(
        Numbering(document, rulebook()).violations(),
        empty(),
        "points inside the subsection were not accepted",
    )


def test_rejects_point_that_skips_a_level() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    document.add_paragraph("1.1.1. Программа прошивает контроллер.")
    assert_that(
        str(Numbering(document, rulebook()).violations()[0]),
        contains_string("пункт «1.1.1. Программа прошивает контроллер.»"),
        "point that skips a level was not rejected",
    )


def test_rejects_section_numbered_like_subsection() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    document.add_paragraph("1.1. ПОДКЛЮЧЕНИЕ", style="Heading 1")
    assert_that(
        str(Numbering(document, rulebook()).violations()[0]),
        contains_string("номер раздела 1.1 состоит не из одного числа"),
        "section numbered like a subsection was not rejected",
    )


def test_rejects_section_number_without_dot() -> None:
    document = Document()
    document.add_paragraph("1 ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    assert_that(
        str(Numbering(document, rulebook()).violations()[0]),
        contains_string("после номера раздела 1 нет точки"),
        "section number without a dot was not rejected",
    )


def test_restarts_after_unnumbered_title() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    document.add_paragraph("2. ОПИСАНИЕ", style="Heading 1")
    document.add_paragraph("ПРИЛОЖЕНИЕ 1", style="Heading 1")
    document.add_paragraph("1. ПОДКЛЮЧЕНИЕ ПЛАТЫ", style="Heading 1")
    assert_that(
        Numbering(document, rulebook()).violations(),
        empty(),
        "numbering did not restart after an unnumbered title",
    )


def test_skips_enumerations() -> None:
    document = Document()
    document.add_paragraph("3. НАСТРОЙКА", style="Heading 1")
    document.add_paragraph("1. Подключить контроллер;")
    document.add_paragraph("2. Открыть программу;")
    document.add_paragraph("1. Type: Serial;")
    assert_that(
        Numbering(document, rulebook()).violations(),
        has_length(1),
        "enumerations were not skipped",
    )


def test_skips_decimal_numbers_in_text() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    document.add_paragraph("3.3 вольта подают на вывод питания.")
    assert_that(
        Numbering(document, rulebook()).violations(),
        empty(),
        "decimal number in the text was not skipped",
    )
