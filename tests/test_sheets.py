from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from hamcrest import assert_that, contains_string, empty, has_length

from gost_espd_checker.sheets import Sheets
from tests.rulebook import rulebook


def test_accepts_first_section_of_document() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    assert_that(
        Sheets(document, rulebook()).violations(),
        empty(),
        "first section of the document was not accepted",
    )


def test_warns_about_section_right_after_text() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    document.add_paragraph("Программа прошивает контроллер.")
    document.add_paragraph("2. ОПИСАНИЕ ПРОГРАММЫ", style="Heading 1")
    assert_that(
        str(Sheets(document, rulebook()).violations()[0]),
        contains_string(
            "заголовок «2. ОПИСАНИЕ ПРОГРАММЫ»: предупреждение ESPD-106-13:"
            " перед разделом нет перехода на новый лист"
        ),
        "section right after text did not get a warning",
    )


def test_accepts_section_after_page_break_at_end_of_text() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    document.add_paragraph(
        "Программа прошивает контроллер."
    ).add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("2. ОПИСАНИЕ ПРОГРАММЫ", style="Heading 1")
    assert_that(
        Sheets(document, rulebook()).violations(),
        empty(),
        "section after a page break at the end of the text was not accepted",
    )


def test_accepts_section_after_page_break_and_empty_paragraphs() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    document.add_paragraph("Программа прошивает контроллер.")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph()
    document.add_paragraph("2. ОПИСАНИЕ ПРОГРАММЫ", style="Heading 1")
    assert_that(
        Sheets(document, rulebook()).violations(),
        empty(),
        "section after a page break and empty paragraphs was not accepted",
    )


def test_accepts_page_break_inside_title_before_its_text() -> None:
    document = Document()
    document.add_paragraph("Программа прошивает контроллер.")
    title = document.add_paragraph(style="Heading 1")
    title.add_run().add_break(WD_BREAK.PAGE)
    title.add_run("3. НАСТРОЙКА ПРОГРАММЫ")
    assert_that(
        Sheets(document, rulebook()).violations(),
        empty(),
        "page break inside the title before its text was not accepted",
    )


def test_warns_when_page_break_follows_text_of_title() -> None:
    document = Document()
    document.add_paragraph("Программа прошивает контроллер.")
    title = document.add_paragraph("3. НАСТРОЙКА ПРОГРАММЫ", style="Heading 1")
    title.add_run().add_break(WD_BREAK.PAGE)
    assert_that(
        Sheets(document, rulebook()).violations(),
        has_length(1),
        "page break after the text of the title did not get a warning",
    )


def test_accepts_page_break_before_set_on_title() -> None:
    document = Document()
    document.add_paragraph("Программа прошивает контроллер.")
    document.add_paragraph(
        "4. ПРОВЕРКА ПРОГРАММЫ", style="Heading 1"
    ).paragraph_format.page_break_before = True
    assert_that(
        Sheets(document, rulebook()).violations(),
        empty(),
        "page break before set on the title was not accepted",
    )


def test_accepts_page_break_before_set_by_style() -> None:
    document = Document()
    document.styles["Heading 1"].paragraph_format.page_break_before = True
    document.add_paragraph("Программа прошивает контроллер.")
    document.add_paragraph("4. ПРОВЕРКА ПРОГРАММЫ", style="Heading 1")
    assert_that(
        Sheets(document, rulebook()).violations(),
        empty(),
        "page break before set by the style was not accepted",
    )


def test_accepts_section_after_break_to_next_page() -> None:
    document = Document()
    document.add_paragraph("Программа прошивает контроллер.")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("5. ДОПОЛНИТЕЛЬНЫЕ ВОЗМОЖНОСТИ", style="Heading 1")
    assert_that(
        Sheets(document, rulebook()).violations(),
        empty(),
        "section after a break to the next page was not accepted",
    )


def test_warns_about_section_after_continuous_break() -> None:
    document = Document()
    document.add_paragraph("Программа прошивает контроллер.")
    document.add_section(WD_SECTION.CONTINUOUS)
    document.add_paragraph("5. ДОПОЛНИТЕЛЬНЫЕ ВОЗМОЖНОСТИ", style="Heading 1")
    assert_that(
        Sheets(document, rulebook()).violations(),
        has_length(1),
        "section after a continuous break did not get a warning",
    )


def test_ignores_subsections_right_after_text() -> None:
    document = Document()
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ", style="Heading 1")
    document.add_paragraph("Программа прошивает контроллер.")
    document.add_paragraph("1.1. Назначение программы", style="Heading 2")
    assert_that(
        Sheets(document, rulebook()).violations(),
        empty(),
        "subsection right after text was not ignored",
    )
