from docx import Document
from docx.enum.text import WD_BREAK
from hamcrest import assert_that, contains_exactly, empty

from gost_espd_checker.title_page import TitlePage


def test_skips_approval_sheet() -> None:
    document = Document()
    document.add_paragraph("ЛИСТ УТВЕРЖДЕНИЯ").add_run().add_break(
        WD_BREAK.PAGE
    )
    document.add_paragraph("Утвержден")
    document.add_paragraph(" ")
    document.add_paragraph("RU.17701729.00512-02 33 01-ЛУ")
    assert_that(
        [line.text for line in TitlePage(document)],
        contains_exactly("Утвержден", "RU.17701729.00512-02 33 01-ЛУ"),
        "title page did not skip the approval sheet and empty lines",
    )


def test_gives_nothing_for_empty_document() -> None:
    assert_that(
        list(TitlePage(Document())),
        empty(),
        "empty document did not give an empty title page",
    )
