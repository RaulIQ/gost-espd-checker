from docx import Document
from hamcrest import assert_that, contains_exactly, empty

from gost_espd_checker.illustrations import Illustrations
from tests.conforming import endorsed
from tests.drawing import pictured


def test_finds_number_under_illustration() -> None:
    document = endorsed(Document())
    document.add_paragraph("Схема обмена")
    pictured(document)
    document.add_paragraph("Стрелки показывают направление.")
    document.add_paragraph("Рис. 4")
    assert_that(
        [(caption, above) for _, caption, above, _ in Illustrations(document)],
        contains_exactly(("Рис. 4", False)),
        "number under the illustration was not found",
    )


def test_finds_number_over_illustration() -> None:
    document = endorsed(Document())
    document.add_paragraph("Рис. 2")
    pictured(document)
    assert_that(
        [(caption, above) for _, caption, above, _ in Illustrations(document)],
        contains_exactly(("Рис. 2", True)),
        "number over the illustration was not found",
    )


def test_gives_number_to_one_illustration_only() -> None:
    document = endorsed(Document())
    pictured(document)
    document.add_paragraph("Рис. 1")
    pictured(document)
    document.add_paragraph("Рис. 2")
    assert_that(
        [caption for _, caption, _, _ in Illustrations(document)],
        contains_exactly("Рис. 1", "Рис. 2"),
        "number was not kept to its own illustration",
    )


def test_counts_appendices() -> None:
    document = endorsed(Document())
    pictured(document)
    document.add_paragraph("ПРИЛОЖЕНИЕ 1")
    pictured(document)
    assert_that(
        [part for _, _, _, part in Illustrations(document)],
        contains_exactly(0, 1),
        "appendices were not counted",
    )


def test_ignores_text_boxes() -> None:
    document = endorsed(Document())
    pictured(document, "wordprocessingShape")
    assert_that(
        list(Illustrations(document)),
        empty(),
        "text box was not left out",
    )
