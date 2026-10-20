from docx import Document
from docx.shared import Mm
from hamcrest import assert_that, contains_string, empty

from gost_espd_checker.subsection_title import SubsectionTitle
from tests.rulebook import rulebook


def test_accepts_indented_title_with_first_capital() -> None:
    paragraph = Document().add_paragraph("1.1. Назначение программы")
    paragraph.paragraph_format.first_line_indent = Mm(12.5)
    assert_that(
        SubsectionTitle(paragraph, "заголовок", rulebook()).violations(),
        empty(),
        "indented title with a first capital was not accepted",
    )


def test_accepts_abbreviation_inside_title() -> None:
    paragraph = Document().add_paragraph("1.3. Модуль BMS на ESP32")
    paragraph.paragraph_format.first_line_indent = Mm(10)
    assert_that(
        SubsectionTitle(paragraph, "заголовок", rulebook()).violations(),
        empty(),
        "abbreviation inside the subsection title was not accepted",
    )


def test_rejects_title_without_indent() -> None:
    assert_that(
        str(
            SubsectionTitle(
                Document().add_paragraph("2.2. Связи между частями"),
                "заголовок «2.2. Связи между частями»",
                rulebook(),
            ).violations()[0]
        ),
        contains_string(
            "заголовок «2.2. Связи между частями»: ошибка ESPD-106-10:"
            " заголовок подраздела записан без абзацного отступа"
        ),
        "subsection title without an indent was not rejected",
    )


def test_rejects_lowercase_first_letter() -> None:
    paragraph = Document().add_paragraph("3.1. установка программы")
    paragraph.paragraph_format.first_line_indent = Mm(12.5)
    assert_that(
        str(
            SubsectionTitle(paragraph, "заголовок", rulebook()).violations()[0]
        ),
        contains_string(
            "ESPD-106-10: заголовок подраздела записан не строчными буквами,"
            " кроме первой прописной"
        ),
        "lowercase first letter of the subsection title was not rejected",
    )


def test_rejects_title_in_capitals() -> None:
    paragraph = Document().add_paragraph("3.2. НАСТРОЙКИ ПРОГРАММЫ")
    paragraph.paragraph_format.first_line_indent = Mm(12.5)
    assert_that(
        str(
            SubsectionTitle(paragraph, "заголовок", rulebook()).violations()[0]
        ),
        contains_string(
            "ESPD-106-10: заголовок подраздела записан не строчными"
        ),
        "subsection title in capitals was not rejected",
    )


def test_rejects_title_shown_in_capitals() -> None:
    paragraph = Document().add_paragraph()
    paragraph.add_run("4.1. Условия проверки").font.all_caps = True
    paragraph.paragraph_format.first_line_indent = Mm(12.5)
    assert_that(
        str(
            SubsectionTitle(paragraph, "заголовок", rulebook()).violations()[0]
        ),
        contains_string(
            "ESPD-106-10: заголовок подраздела записан не строчными"
        ),
        "subsection title shown in capitals was not rejected",
    )
