from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from hamcrest import assert_that, contains_string, empty

from gost_espd_checker.section_title import SectionTitle
from tests.rulebook import rulebook


def test_accepts_centered_capitals() -> None:
    paragraph = Document().add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ ПРОГРАММЫ")
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    assert_that(
        SectionTitle(paragraph, "заголовок", rulebook()).violations(),
        empty(),
        "centered capitals were not accepted",
    )


def test_accepts_lowercase_shown_in_capitals() -> None:
    paragraph = Document().add_paragraph()
    paragraph.add_run("2. Структура программы").font.all_caps = True
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    assert_that(
        SectionTitle(paragraph, "заголовок", rulebook()).violations(),
        empty(),
        "lowercase shown in capitals was not accepted",
    )


def test_accepts_abbreviation_with_digits() -> None:
    paragraph = Document().add_paragraph("3. ПРОШИВКА ESP32-S3")
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    assert_that(
        SectionTitle(paragraph, "заголовок", rulebook()).violations(),
        empty(),
        "abbreviation with digits was not accepted",
    )


def test_rejects_lowercase_letters() -> None:
    paragraph = Document().add_paragraph("4. Проверка программы")
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    assert_that(
        str(
            SectionTitle(
                paragraph, "заголовок «4. Проверка программы»", rulebook()
            ).violations()[0]
        ),
        contains_string(
            "заголовок «4. Проверка программы»: ошибка ESPD-106-08:"
            " заголовок раздела написан не прописными буквами"
        ),
        "lowercase letters in the section title were not rejected",
    )


def test_rejects_left_alignment() -> None:
    assert_that(
        str(
            SectionTitle(
                Document().add_paragraph("5. ДОПОЛНИТЕЛЬНЫЕ ВОЗМОЖНОСТИ"),
                "заголовок",
                rulebook(),
            ).violations()[0]
        ),
        contains_string("ESPD-106-09: заголовок раздела не выровнен по центру"),
        "section title aligned left was not rejected",
    )
