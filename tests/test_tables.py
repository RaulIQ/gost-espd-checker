from docx import Document
from hamcrest import assert_that, contains_exactly

from gost_espd_checker.tables import Tables
from tests.conforming import endorsed
from tests.tabled import tabled


def test_takes_caption_right_over_table() -> None:
    document = endorsed(Document())
    tabled(document, "Таблица 3 — Устройства на шине", ("Устройство", "Адрес"))
    assert_that(
        [caption for _, caption, _ in Tables(document)],
        contains_exactly("Таблица 3 — Устройства на шине"),
        "caption right over the table was not taken",
    )


def test_leaves_caption_empty_without_word_table() -> None:
    document = endorsed(Document())
    tabled(document, "Режимы работы платы", ("Режим", "Признак"))
    assert_that(
        [caption for _, caption, _ in Tables(document)],
        contains_exactly(""),
        "caption was not left empty without the word «Таблица»",
    )


def test_skips_sheet_of_changes() -> None:
    document = endorsed(Document())
    tabled(document, "Таблица 1 — Режимы", ("Режим", "Признак"))
    document.add_paragraph("ЛИСТ РЕГИСТРАЦИИ ИЗМЕНЕНИЙ")
    tabled(document, "", ("Изм.", "Дата", "Подпись"))
    assert_that(
        [caption for _, caption, _ in Tables(document)],
        contains_exactly("Таблица 1 — Режимы"),
        "sheet of changes was not skipped",
    )


def test_names_appendix_of_table() -> None:
    document = endorsed(Document())
    tabled(document, "Таблица 1 — Режимы", ("Режим", "Признак"))
    document.add_paragraph("ПРИЛОЖЕНИЕ 2")
    tabled(document, "Таблица 2.1 — Коды", ("Код", "Значение"))
    assert_that(
        [appendix for _, _, appendix in Tables(document)],
        contains_exactly("", "2"),
        "appendix of the table was not named",
    )


def test_skips_list_under_its_heading() -> None:
    document = endorsed(Document())
    tabled(document, "Таблица 1 — Режимы", ("Режим", "Признак"))
    document.add_paragraph("ПЕРЕЧЕНЬ СОКРАЩЕНИЙ")
    tabled(document, "", ("Сокращение", "Пояснение"))
    assert_that(
        [caption for _, caption, _ in Tables(document)],
        contains_exactly("Таблица 1 — Режимы"),
        "list under its heading was not skipped",
    )
