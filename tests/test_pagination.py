import secrets

from docx import Document
from docx.document import Document as Word
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from hamcrest import assert_that, contains_string, empty, has_length

from gost_espd_checker.pagination import Pagination
from tests.rulebook import rulebook


def restarted(document: Word, index: int, start: int) -> Word:
    document.element.body.xpath("./w:p/w:pPr/w:sectPr | ./w:sectPr")[
        index
    ].append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="{start}"/>'))
    return document


def test_accepts_start_of_first_section() -> None:
    document = Document()
    document.add_paragraph("Программа прошивает контроллер.")
    assert_that(
        Pagination(restarted(document, 0, 1), rulebook()).violations(),
        empty(),
        "start of the first section was not accepted",
    )


def test_accepts_sections_that_continue_numbering() -> None:
    document = Document()
    document.add_paragraph("Программа прошивает контроллер.")
    document.add_section()
    document.add_paragraph("Загрузчик проверяет образ.")
    assert_that(
        Pagination(document, rulebook()).violations(),
        empty(),
        "sections that continue the numbering were not accepted",
    )


def test_rejects_second_section_that_restarts_numbering() -> None:
    start = secrets.randbelow(20) + 1
    document = Document()
    document.add_paragraph("Программа прошивает контроллер.")
    document.add_section()
    document.add_paragraph("Загрузчик проверяет образ.")
    assert_that(
        str(
            Pagination(restarted(document, 1, start), rulebook()).violations()[
                0
            ]
        ),
        contains_string(
            f"раздел 2: ошибка ESPD-106-05:"
            f" нумерация страниц начинается заново от {start}"
        ),
        "second section that restarts the numbering was not rejected",
    )


def test_rejects_every_section_that_restarts_numbering() -> None:
    document = Document()
    document.add_paragraph("Программа прошивает контроллер.")
    document.add_section()
    document.add_section()
    assert_that(
        Pagination(
            restarted(restarted(document, 1, 1), 2, 1), rulebook()
        ).violations(),
        has_length(2),
        "every section that restarts the numbering was not rejected",
    )
