from docx.document import Document
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.section import Section
from docx.shared import Mm


def conforming(section: Section) -> Section:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25)
    section.left_margin = Mm(20)
    section.right_margin = Mm(10)
    section.bottom_margin = Mm(15)
    section.header.is_linked_to_previous = False
    headed(
        section,
        '<w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>',
    )
    return headed(section, "<w:r><w:t>RU.17701729.00001-01 34 01</w:t></w:r>")


def headed(section: Section, content: str) -> Section:
    section.header.part.element.append(
        parse_xml(f"<w:p {nsdecls('w')}>{content}</w:p>")
    )
    return section


def endorsed(document: Document) -> Document:
    document.add_paragraph("Утвержден")
    document.add_paragraph("RU.17701729.00001-01 32 01-ЛУ").add_run().add_break(
        WD_BREAK.PAGE
    )
    return document
