from docx import Document
from hamcrest import assert_that, contains_exactly, empty

from gost_espd_checker.table_heads import TableHeads
from tests.conforming import endorsed
from tests.rulebook import rulebook
from tests.tabled import tabled


def test_accepts_heads_in_capital_without_period() -> None:
    document = endorsed(Document())
    tabled(
        document,
        "Таблица 1 — Устройства",
        ("Устройство", "Адрес"),
        ("Зарядное устройство", "0x6B"),
        ("bms-esp32s3-full.bin", "0x0"),
    )
    assert_that(
        TableHeads(document, rulebook()).violations(),
        empty(),
        "heads in a capital letter without a period were not accepted",
    )


def test_rejects_column_head_in_small_letter() -> None:
    document = endorsed(Document())
    tabled(document, "Таблица 1 — Устройства", ("Устройство", "адрес"))
    assert_that(
        [
            str(violation)
            for violation in TableHeads(document, rulebook()).violations()
        ],
        contains_exactly(
            "таблица «Таблица 1 — Устройства»: ошибка ESPD-015-03: первая"
            " буква заголовка «адрес» строчная (ГОСТ 1.5-2001, п. 4.5.4)"
        ),
        "column head in a small letter was not rejected",
    )


def test_rejects_row_head_with_period() -> None:
    document = endorsed(Document())
    tabled(
        document,
        "Таблица 2 — Режимы",
        ("Режим", "Признак"),
        ("Ожидание.", "Зелёный индикатор"),
    )
    assert_that(
        [
            str(violation)
            for violation in TableHeads(document, rulebook()).violations()
        ],
        contains_exactly(
            "таблица «Таблица 2 — Режимы»: ошибка ESPD-015-03: в конце"
            " заголовка «Ожидание.» стоит точка (ГОСТ 1.5-2001, п. 4.5.4)"
        ),
        "row head with a period was not rejected",
    )
