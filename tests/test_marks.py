from docx import Document
from docx.enum.text import WD_BREAK
from hamcrest import assert_that, contains_exactly

from gost_espd_checker.marks import Marks


def test_lists_page_breaks_and_text_in_order() -> None:
    paragraph = Document().add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run("Сообщения оператору")
    paragraph.add_run("  ")
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    assert_that(
        ["разрыв" if mark else "текст" for mark in Marks(paragraph)],
        contains_exactly("разрыв", "текст", "разрыв"),
        "page breaks and text did not come in order",
    )
