from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Mm
from hamcrest import assert_that, contains_string, empty

from gost_espd_checker.titles import Titles
from tests.rulebook import rulebook


def test_accepts_document_with_proper_titles() -> None:
    document = Document()
    document.add_paragraph(
        "1. ОБЩИЕ СВЕДЕНИЯ"
    ).alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    document.add_paragraph(
        "1.1. Назначение программы"
    ).paragraph_format.first_line_indent = Mm(12.5)
    document.add_paragraph("Программа прошивает контроллер.")
    assert_that(
        Titles(document, rulebook()).violations(),
        empty(),
        "document with proper titles was not accepted",
    )


def test_finds_section_title_by_number() -> None:
    document = Document()
    document.add_paragraph("2. Структура программы")
    assert_that(
        "\n".join(
            str(item) for item in Titles(document, rulebook()).violations()
        ),
        contains_string(
            "заголовок «2. Структура программы»: ошибка ESPD-106-08"
        ),
        "section title was not found by its number",
    )


def test_finds_section_title_by_style() -> None:
    document = Document()
    document.add_paragraph("Структура программы", style="Heading 1")
    assert_that(
        "\n".join(
            str(item) for item in Titles(document, rulebook()).violations()
        ),
        contains_string("ESPD-106-09"),
        "section title was not found by its style",
    )


def test_finds_subsection_title_by_number() -> None:
    document = Document()
    document.add_paragraph("2.3. Связи между устройствами")
    assert_that(
        "\n".join(
            str(item) for item in Titles(document, rulebook()).violations()
        ),
        contains_string("ESPD-106-10: заголовок подраздела записан без"),
        "subsection title was not found by its number",
    )


def test_prefers_style_to_number() -> None:
    document = Document()
    document.add_paragraph("6. Сообщения журнала", style="Heading 2")
    assert_that(
        "\n".join(
            str(item) for item in Titles(document, rulebook()).violations()
        ),
        contains_string("ESPD-106-10: заголовок подраздела записан без"),
        "style of the title was not preferred to its number",
    )


def test_skips_lines_of_contents() -> None:
    document = Document()
    document.add_paragraph("1.1. Назначение программы\t4")
    assert_that(
        Titles(document, rulebook()).violations(),
        empty(),
        "line of the contents was not skipped",
    )


def test_skips_deeper_headings() -> None:
    document = Document()
    document.add_paragraph("Порядок прошивки.", style="Heading 3")
    assert_that(
        Titles(document, rulebook()).violations(),
        empty(),
        "deeper heading was not skipped",
    )


def test_skips_plain_paragraphs() -> None:
    document = Document()
    document.add_paragraph("Программа записывает образ во флеш-память.")
    assert_that(
        Titles(document, rulebook()).violations(),
        empty(),
        "plain paragraph was not skipped",
    )


def test_reads_automatic_hyphenation_of_document() -> None:
    document = Document()
    document.settings.element.append(
        parse_xml(f"<w:autoHyphenation {nsdecls('w')}/>")
    )
    document.add_paragraph(
        "3. НАСТРОЙКА"
    ).alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    assert_that(
        "\n".join(
            str(item) for item in Titles(document, rulebook()).violations()
        ),
        contains_string("ESPD-106-11: в заголовке разрешены автоматические"),
        "automatic hyphenation of the document was not read",
    )
