from docx import Document
from docx.shared import Mm
from hamcrest import assert_that, contains_string, empty

from gost_espd_checker.layout import Layout
from tests.conforming import conforming
from tests.rulebook import rulebook


def test_accepts_document_of_conforming_sections() -> None:
    document = Document()
    conforming(document.sections[0])
    conforming(document.add_section())
    assert_that(
        Layout(document, rulebook()).violations(),
        empty(),
        "document of conforming sections was not accepted",
    )


def test_reports_wrong_sheet_of_second_section() -> None:
    document = Document()
    conforming(document.sections[0])
    conforming(document.add_section()).page_height = Mm(280)
    assert_that(
        "\n".join(
            str(item) for item in Layout(document, rulebook()).violations()
        ),
        contains_string("раздел 2: ошибка ESPD-106-01: лист"),
        "wrong sheet of the second section was not reported",
    )


def test_reports_wrong_margin_of_first_section() -> None:
    document = Document()
    conforming(document.sections[0]).top_margin = Mm(20)
    assert_that(
        "\n".join(
            str(item) for item in Layout(document, rulebook()).violations()
        ),
        contains_string("раздел 1: ошибка ESPD-106-02: верхнее поле"),
        "wrong margin of the first section was not reported",
    )


def test_reports_missing_page_number_of_second_section() -> None:
    document = Document()
    conforming(document.sections[0])
    second = document.add_section()
    second.header.is_linked_to_previous = False
    assert_that(
        "\n".join(
            str(item) for item in Layout(document, rulebook()).violations()
        ),
        contains_string("раздел 2: ошибка ESPD-106-03: в верхнем колонтитуле"),
        "missing page number of the second section was not reported",
    )
