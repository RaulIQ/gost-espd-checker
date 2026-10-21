from docx import Document
from docx.enum.text import WD_BREAK
from hamcrest import assert_that, contains_exactly

from gost_espd_checker.lines import Lines


def test_numbers_pages_of_lines_with_text() -> None:
    document = Document()
    document.add_paragraph("АННОТАЦИЯ").add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("   ")
    document.add_paragraph("СОДЕРЖАНИЕ")
    assert_that(
        [(page, line.text) for page, line in Lines(document)],
        contains_exactly((0, "АННОТАЦИЯ"), (1, "СОДЕРЖАНИЕ")),
        "lines with text did not get the numbers of their pages",
    )
