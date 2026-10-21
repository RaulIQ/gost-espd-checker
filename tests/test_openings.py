from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from hamcrest import assert_that, contains_exactly

from gost_espd_checker.openings import Openings


def test_sees_new_page_after_page_break() -> None:
    document = Document()
    document.add_paragraph("Аннотация").add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("1. ОБЩИЕ СВЕДЕНИЯ")
    assert_that(
        ["лист" if fresh else "текст" for fresh in Openings(document)],
        contains_exactly("лист", "лист"),
        "paragraph after the page break did not start a new page",
    )


def test_sees_text_before_paragraph() -> None:
    document = Document()
    document.add_paragraph("Программа работает автономно.")
    document.add_paragraph("2. НАЗНАЧЕНИЕ")
    assert_that(
        ["лист" if fresh else "текст" for fresh in Openings(document)],
        contains_exactly("лист", "текст"),
        "text before the paragraph was not seen",
    )


def test_sees_new_page_after_section_break() -> None:
    document = Document()
    document.add_paragraph("Листов 9")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("АННОТАЦИЯ")
    assert_that(
        ["лист" if fresh else "текст" for fresh in Openings(document)],
        contains_exactly("лист", "текст", "лист"),
        "paragraph after the section break did not start a new page",
    )
