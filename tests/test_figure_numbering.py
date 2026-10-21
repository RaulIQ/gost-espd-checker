from docx import Document
from docx.document import Document as Word
from hamcrest import assert_that, contains_exactly, empty

from gost_espd_checker.figure_numbering import FigureNumbering
from tests.conforming import endorsed
from tests.drawing import pictured
from tests.rulebook import rulebook


def drawn(*captions: str) -> Word:
    document = endorsed(Document())
    for caption in captions:
        if caption.startswith("ПРИЛОЖЕНИЕ"):
            document.add_paragraph(caption)
            continue
        pictured(document)
        document.add_paragraph(caption)
    return document


def test_accepts_numbers_in_order() -> None:
    assert_that(
        FigureNumbering(
            drawn("Рис. 1", "Рис. 2", "Рис. 3"), rulebook()
        ).violations(),
        empty(),
        "numbers in order were not accepted",
    )


def test_accepts_single_illustration_without_number() -> None:
    assert_that(
        FigureNumbering(drawn("Окно программы"), rulebook()).violations(),
        empty(),
        "single illustration without a number was not accepted",
    )


def test_rejects_illustration_without_number() -> None:
    assert_that(
        [
            str(violation)
            for violation in FigureNumbering(
                drawn("Рис. 1", "Окно настройки"), rulebook()
            ).violations()
        ],
        contains_exactly(
            "иллюстрация 2 по порядку: ошибка ESPD-106-18: иллюстрация не"
            " пронумерована (ГОСТ 19.106-78, п. 2.3.1)"
        ),
        "illustration without a number was not rejected",
    )


def test_rejects_skipped_number() -> None:
    assert_that(
        [
            str(violation)
            for violation in FigureNumbering(
                drawn("Рис. 1", "Рис. 3"), rulebook()
            ).violations()
        ],
        contains_exactly(
            "иллюстрация «Рис. 3»: ошибка ESPD-106-18: номер 3 вместо 2"
            " (ГОСТ 19.106-78, п. 2.3.1)"
        ),
        "skipped number was not rejected",
    )


def test_restarts_numbers_in_appendix() -> None:
    assert_that(
        FigureNumbering(
            drawn("Рис. 1", "Рис. 2", "ПРИЛОЖЕНИЕ 1", "Рис. 1"), rulebook()
        ).violations(),
        empty(),
        "numbers did not restart in the appendix",
    )
